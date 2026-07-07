#!/usr/bin/env python3
"""
Process the 'USP Evolution - 172M' sheet.

Layout:
  - Columns A-D: A=Site-Template, B=Network Element, C=CNF/VNF, D=Subs
  - Row 23, from column E onward: site names. Each site header spans
    multiple columns (normally 6); width is derived from the distance
    between consecutive non-empty cells on row 23.
  - Row 25 onward: per-function rows. Each site block contains formulas
    referencing 'Detailed Dimensioning- 172M', e.g.
        ='Detailed Dimensioning- 172M'!$K75*2
    The DM row (75) is captured as "DMrow".

Output: <input-basename>-usp-evolution.json
"""

import sys
import re
import json
from pathlib import Path
from collections import OrderedDict
import openpyxl

SHEET_NAME     = "USP Evolution - 172M"
DM_SHEET_NAME  = "Detailed Dimensioning- 172M"

LOCATION_ROW                    = 23
LOCATION_START_COL              = 5      # E
LOCATION_STOP_AFTER_EMPTY_COLS  = 20

START_ROW           = 25
DEFAULT_SITE_WIDTH  = 6

COL_SITE_TEMPLATE   = 1  # A
COL_NETWORK_ELEMENT = 2  # B
COL_CNF_VNF         = 3  # C
COL_SUBS            = 4  # D

SKIP_FROM             = "3rd Party Onboarding"
SKIP_UNTIL_INCLUSIVE  = "Infrastructure"
STOP_CONTAINS         = "Additional (Nokia)"

# Match  'Detailed Dimensioning- 172M'!$K75   or   Detailed Dimensioning- 172M!K75
DM_ROW_RE = re.compile(
    r"'?" + re.escape(DM_SHEET_NAME) + r"'?!\$?[A-Za-z]+\$?(\d+)"
)


# ---------- helpers ----------------------------------------------------------

def normalize_cell_text(v) -> str:
    s = str(v).strip().replace("\r", " ").replace("\n", " ")
    return " ".join(s.split())


def cell_to_string(v) -> str:
    return "" if v is None else normalize_cell_text(v)


def extract_dm_row(formula):
    """Return the DM sheet row number referenced in a formula, or None."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    m = DM_ROW_RE.search(formula)
    return int(m.group(1)) if m else None


def read_sites(ws):
    """Read site headers on row 23 starting at column E.
       Return list of (start_col, end_col, site_name)."""
    raw = []
    empty_streak = 0
    col = LOCATION_START_COL
    while True:
        s = cell_to_string(ws.cell(row=LOCATION_ROW, column=col).value)
        if s:
            raw.append((col, s))
            empty_streak = 0
        else:
            empty_streak += 1
            if empty_streak >= LOCATION_STOP_AFTER_EMPTY_COLS:
                break
        col += 1

    sites = []
    for i, (start_col, name) in enumerate(raw):
        if i + 1 < len(raw):
            end_col = raw[i + 1][0] - 1
        else:
            end_col = start_col + DEFAULT_SITE_WIDTH - 1
        sites.append((start_col, end_col, name))
    return sites


# ---------- main -------------------------------------------------------------

def main():
    if len(sys.argv) != 2:
        print(f"Usage: {Path(sys.argv[0]).name} <workbook.xlsx>", file=sys.stderr)
        sys.exit(2)

    in_path = Path(sys.argv[1])

    # Only need raw formulas for DMrow extraction
    wb = openpyxl.load_workbook(in_path, data_only=False)
    if SHEET_NAME not in wb.sheetnames:
        print(f"Sheet '{SHEET_NAME}' not found in {in_path}", file=sys.stderr)
        sys.exit(1)
    ws = wb[SHEET_NAME]

    sites = read_sites(ws)
    if not sites:
        print(f"No sites found on row {LOCATION_ROW}", file=sys.stderr)
        sys.exit(1)

    # Report unusual widths so anything odd is visible
    for start, end, name in sites:
        w = end - start + 1
        if w != DEFAULT_SITE_WIDTH:
            print(f"[info] site '{name}' spans {w} columns "
                  f"({openpyxl.utils.get_column_letter(start)}"
                  f"-{openpyxl.utils.get_column_letter(end)})", file=sys.stderr)

    site_data = OrderedDict()   # site -> OrderedDict(template -> [ne dicts])
    for _, _, sname in sites:
        site_data[sname] = OrderedDict()

    row = START_ROW
    max_row = ws.max_row
    current_template = ""
    skipping = False

    while row <= max_row:
        tmpl = cell_to_string(ws.cell(row=row, column=COL_SITE_TEMPLATE).value)
        ne   = cell_to_string(ws.cell(row=row, column=COL_NETWORK_ELEMENT).value)
        cnf  = cell_to_string(ws.cell(row=row, column=COL_CNF_VNF).value)
        subs = cell_to_string(ws.cell(row=row, column=COL_SUBS).value)

        if tmpl:
            current_template = tmpl

        combined = f"{tmpl} {ne}".strip()
        if STOP_CONTAINS and STOP_CONTAINS in combined:
            break

        if SKIP_FROM and (tmpl == SKIP_FROM or ne == SKIP_FROM):
            skipping = True
        if skipping:
            if SKIP_UNTIL_INCLUSIVE and (
                tmpl == SKIP_UNTIL_INCLUSIVE or ne == SKIP_UNTIL_INCLUSIVE
            ):
                skipping = False
            row += 1
            continue

        if not ne:
            row += 1
            continue

        for start_col, end_col, sname in sites:
            dm_rows_seen = []
            for c in range(start_col, end_col + 1):
                dmr = extract_dm_row(ws.cell(row=row, column=c).value)
                if dmr is not None and dmr not in dm_rows_seen:
                    dm_rows_seen.append(dmr)

            if len(dm_rows_seen) > 1:
                print(f"[warn] row {row} site '{sname}' has multiple DM rows "
                      f"{dm_rows_seen}; using {dm_rows_seen[0]}", file=sys.stderr)
            dm_row = dm_rows_seen[0] if dm_rows_seen else None

            tkey = current_template
            if tkey not in site_data[sname]:
                site_data[sname][tkey] = []
            site_data[sname][tkey].append({
                "network_element": ne,
                "CNF_VNF": cnf,
                "Subs": subs,
                "line": row,
                "DMrow": dm_row,
            })

        row += 1

    # Build final JSON
    output = []
    for sname, templates in site_data.items():
        entry = {"site": sname, "Site-Templates": []}
        for tname, nes in templates.items():
            entry["Site-Templates"].append({
                "site-template": tname,
                "network_elements": nes,
            })
        output.append(entry)

    out_path = in_path.with_name(f"{in_path.stem}-usp-evolution.json")
    with out_path.open("w", encoding="utf-8") as f:
        json.dump(output, f, indent=2)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()