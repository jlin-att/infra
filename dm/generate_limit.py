#!/usr/bin/env python3
"""
Extract CNF resource limits from an XLSX file and output to a JSON file.
Version is parsed from the input filename (e.g., v18.1 -> "18.1").
Output filename: <basename>-limit.json
"""

import sys
import os
import re
import json
import argparse
import openpyxl
from openpyxl.utils import get_column_letter


def extract_version(filename):
    """Extract version number from filename pattern like _v18.1_ or _v2.0_"""
    basename = os.path.basename(filename)
    match = re.search(r"_v(\d+(?:\.\d+)+)", basename, re.IGNORECASE)
    if match:
        return match.group(1)
    # Fallback: try without underscore prefix
    match = re.search(r"v(\d+(?:\.\d+)+)", basename, re.IGNORECASE)
    if match:
        return match.group(1)
    return "unknown"


def build_output_path(input_path):
    """Build output JSON path: <basename>-limit.json"""
    dirname = os.path.dirname(input_path)
    basename = os.path.splitext(os.path.basename(input_path))[0]
    output_name = f"{basename}-limit.json"
    return os.path.join(dirname, output_name) if dirname else output_name


def main():
    parser = argparse.ArgumentParser(
        description="Extract CNF CPU/Memory limits to JSON."
    )
    parser.add_argument("xlsx", help="Path to the input XLSX file")
    args = parser.parse_args()

    xlsx_path = args.xlsx
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    sheet_name = "CPU Memory Limits"
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[sheet_name]

    # ── Parse version from filename ────────────────────────────────────
    version = extract_version(xlsx_path)

    # ── Locate header row (row 5, 1-indexed) ──────────────────────────
    header_row = 5
    headers = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is not None:
            headers[col] = str(val).strip()

    cnf_col = 1  # Column A

    # Find the FIRST occurrence of each metric
    cpu_col = None
    mem_col = None
    for col in sorted(headers.keys()):
        if headers[col] == "CPU Limits (milli)" and cpu_col is None:
            cpu_col = col
        if headers[col] == "Memory Limits (GiB)" and mem_col is None:
            mem_col = col

    if cpu_col is None:
        print("ERROR: No column with header 'CPU Limits (milli)' found.")
        sys.exit(1)
    if mem_col is None:
        print("ERROR: No column with header 'Memory Limits (GiB)' found.")
        sys.exit(1)

    print(f"  Version detected  : {version}")
    print(f"  CPU Limits column : {get_column_letter(cpu_col)} (col {cpu_col})")
    print(f"  Memory Limits col : {get_column_letter(mem_col)} (col {mem_col})")

    # ── Walk data rows ─────────────────────────────────────────────────
    data_start = header_row + 1
    cnf_list = []

    for row in range(data_start, ws.max_row + 1):
        cnf = ws.cell(row=row, column=cnf_col).value
        if cnf is None or str(cnf).strip() == "":
            continue

        cpu_val = ws.cell(row=row, column=cpu_col).value
        mem_val = ws.cell(row=row, column=mem_col).value

        cnf_list.append({
            "cnf": str(cnf).strip(),
            "cpu_limits_milli": cpu_val,
            "memory_limits_gib": mem_val,
        })

    # ── Build output JSON ──────────────────────────────────────────────
    output = {
        "version": version,
        "source": os.path.basename(xlsx_path),
        "cnf_count": len(cnf_list),
        "cnfs": cnf_list,
    }

    output_path = build_output_path(xlsx_path)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, default=str)

    print(f"\n  CNFs extracted    : {len(cnf_list)}")
    print(f"  Output written to : {output_path}")

    wb.close()


if __name__ == "__main__":
    main()