#!/usr/bin/env python3
"""
Compare "CPU Limits (milli)" and "Memory Limits (GiB)" columns
across a sheet named "CPU Memory Limits" in an XLSX file.
Row 5 is treated as the header row.
"""

import sys
import argparse
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string


def main():
    parser = argparse.ArgumentParser(
        description="Compare CPU/Memory limit columns in an XLSX file."
    )
    parser.add_argument("xlsx", help="Path to the input XLSX file")
    parser.add_argument(
        "--show-lab",
        action="store_true",
        default=False,
        help="Display 'diff in lab only' entries (hidden by default)",
    )
    args = parser.parse_args()

    xlsx_path = args.xlsx
    show_lab = args.show_lab

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    sheet_name = "CPU Memory Limits"
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[sheet_name]

    # ── Column CK threshold (1-based index) ───────────────────────────
    CK_COL = column_index_from_string("CK")  # 89

    # ── Locate header row (row 5, 1-indexed) ──────────────────────────
    header_row = 5
    headers = {}  # col_idx (1-based) -> header text
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is not None:
            headers[col] = str(val).strip()

    # Column A is the CNF identifier
    cnf_col = 1

    # Collect column indices for each metric
    cpu_cols = sorted([c for c, h in headers.items() if h == "CPU Limits (milli)"])
    mem_cols = sorted([c for c, h in headers.items() if h == "Memory Limits (GiB)"])

    if not cpu_cols:
        print("WARNING: No columns with header 'CPU Limits (milli)' found.")
    if not mem_cols:
        print("WARNING: No columns with header 'Memory Limits (GiB)' found.")

    if not cpu_cols and not mem_cols:
        print("Nothing to compare. Exiting.")
        sys.exit(0)

    # Helper: friendly column label
    def col_label(col_idx):
        parent = ws.cell(row=header_row - 1, column=col_idx).value
        letter = get_column_letter(col_idx)
        if parent:
            return f"{str(parent).strip()} (col {letter})"
        return f"col {letter}"

    def report_diff(row, cnf, metric_name, cols):
        """
        Compare values across cols for a given row.
        Returns: (is_match: bool, is_lab_only: bool)
        """
        vals = {c: ws.cell(row=row, column=c).value for c in cols}
        unique = set(vals.values())

        if len(unique) == 1:
            return True, False  # all match

        first_col = cols[0]
        first_val = vals[first_col]

        diff_cols_before_ck = [c for c in cols[1:] if vals[c] != first_val and c < CK_COL]
        diff_cols_ck_or_after = [c for c in cols[1:] if vals[c] != first_val and c >= CK_COL]

        # Only differences are at/after CK → "diff in lab only"
        if not diff_cols_before_ck and diff_cols_ck_or_after:
            if show_lab:
                print(f"\n[DIFF] Row {row} | CNF: {cnf} | {metric_name}")
                print(f"         >> diff in lab only")
            return False, True

        # Differences before CK (may also have lab diffs)
        print(f"\n[DIFF] Row {row} | CNF: {cnf} | {metric_name}")
        print(f"         {col_label(first_col):>40s} = {first_val}")
        for c in diff_cols_before_ck:
            print(f"         {col_label(c):>40s} = {vals[c]}")
        if diff_cols_ck_or_after:
            print(f"         >> also diff in lab only")

        return False, False

    # ── Walk data rows ─────────────────────────────────────────────────
    data_start = header_row + 1
    cpu_match_count = 0
    mem_match_count = 0
    cpu_diff_count = 0
    mem_diff_count = 0
    cpu_lab_only_count = 0
    mem_lab_only_count = 0

    print("=" * 72)
    print(f"  File  : {xlsx_path}")
    print(f"  Sheet : {sheet_name}")
    print(f"  CPU columns found : {len(cpu_cols)}  |  Memory columns found : {len(mem_cols)}")
    print(f"  Lab-only threshold: col CK ({CK_COL}) and beyond")
    print(f"  Show lab-only diffs: {'Yes' if show_lab else 'No (use --show-lab)'}")
    print("=" * 72)

    for row in range(data_start, ws.max_row + 1):
        cnf = ws.cell(row=row, column=cnf_col).value
        if cnf is None or str(cnf).strip() == "":
            continue

        cnf = str(cnf).strip()

        if len(cpu_cols) > 1:
            matched, lab_only = report_diff(row, cnf, "CPU Limits (milli)", cpu_cols)
            if matched:
                cpu_match_count += 1
            elif lab_only:
                cpu_lab_only_count += 1
            else:
                cpu_diff_count += 1

        if len(mem_cols) > 1:
            matched, lab_only = report_diff(row, cnf, "Memory Limits (GiB)", mem_cols)
            if matched:
                mem_match_count += 1
            elif lab_only:
                mem_lab_only_count += 1
            else:
                mem_diff_count += 1

    # ── Summary ────────────────────────────────────────────────────────
    print("\n" + "=" * 72)
    print("  SUMMARY")
    print("=" * 72)
    if len(cpu_cols) > 1:
        print(f"  CPU Limits (milli)  — matched: {cpu_match_count}  |  "
              f"differences: {cpu_diff_count}  |  lab-only diffs: {cpu_lab_only_count}")
    if len(mem_cols) > 1:
        print(f"  Memory Limits (GiB) — matched: {mem_match_count}  |  "
              f"differences: {mem_diff_count}  |  lab-only diffs: {mem_lab_only_count}")
    print("=" * 72)

    wb.close()


if __name__ == "__main__":
    main()