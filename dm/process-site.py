#!/usr/bin/env python3
import sys
import json
from pathlib import Path

import openpyxl

SHEET_NAME = "CNF VNF Counts - 172M"

# Locations are on row 21, starting from column E
LOCATION_ROW = 21
LOCATION_START_COL = 5  # E
LOCATION_STOP_AFTER_EMPTY = 3  # stop after N consecutive empty cells

# Data starts here
START_ROW = 23

# Column indices
COL_SITE_TEMPLATE = 1    # A
COL_NETWORK_ELEMENT = 2  # B
COL_CNF_VNF = 3          # C
COL_SUBS = 4             # D

SKIP_FROM = "3rd Party Onboarding"
SKIP_UNTIL_INCLUSIVE = "Infrastructure"
STOP_CONTAINS = "Additional (Nokia)"


def normalize_cell_text(v: object) -> str:
    s = str(v).strip()
    s = s.replace("\r", " ").replace("\n", " ")
    s = " ".join(s.split())
    return s


def cell_to_string(v: object) -> str:
    if v is None:
        return ""
    return normalize_cell_text(v)


def excel_col_name(n: int) -> str:
    name = ""
    while n:
        n, r = divmod(n - 1, 26)
        name = chr(65 + r) + name
    return name


def to_int_or_zero(v: object) -> int:
    if v is None or isinstance(v, bool):
        return 0
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else 0

    s = cell_to_string(v)
    if not s:
        return 0
    s = s.replace(",", "")
    try:
        f = float(s)
        return int(f) if f.is_integer() else 0
    except ValueError:
        return 0


def main():
    if len(sys.argv) != 2:
        print(f"Usage: {Path(sys.argv[0]).name} <workbook.xlsx>", file=sys.stderr)
        sys.exit(2)

    wb_path = Path(sys.argv[1])
    if not wb_path.exists():
        print(f"File not found: {wb_path}", file=sys.stderr)
        sys.exit(2)

    # Output file: <basename> + "-cnfvnf".json
    out_path = wb_path.with_name(f"{wb_path.stem}-cnfvnf.json")

    wb = openpyxl.load_workbook(wb_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"Sheet not found: {SHEET_NAME}", file=sys.stderr)
        print("Available sheets:", ", ".join(wb.sheetnames), file=sys.stderr)
        sys.exit(2)

    ws = wb[SHEET_NAME]

    # --- Read locations (row 21 starting col E; stop after 3 consecutive empties) ---
    locations = []
    empty_run = 0
    for c in range(LOCATION_START_COL, ws.max_column + 1):
        name = cell_to_string(ws.cell(row=LOCATION_ROW, column=c).value)
        if not name:
            empty_run += 1
            if empty_run >= LOCATION_STOP_AFTER_EMPTY:
                break
            continue

        empty_run = 0
        locations.append({
            "col": excel_col_name(c),
            "col_idx": c,
            "name": name
        })

    # --- Build Site-Templates structure from Col A/B/C/D + per-location columns ---
    in_skip_block = False
    site_templates = []
    current_st = None

    def start_new_site_template(st_name: str):
        nonlocal current_st
        current_st = {
            "site-template": st_name,
            "network_elements": []
        }
        site_templates.append(current_st)

    for r in range(START_ROW, ws.max_row + 1):
        a_txt = cell_to_string(ws.cell(row=r, column=COL_SITE_TEMPLATE).value)
        b_txt = cell_to_string(ws.cell(row=r, column=COL_NETWORK_ELEMENT).value)
        c_txt = cell_to_string(ws.cell(row=r, column=COL_CNF_VNF).value)
        d_txt = cell_to_string(ws.cell(row=r, column=COL_SUBS).value)

        if a_txt and STOP_CONTAINS.lower() in a_txt.lower():
            break

        if not in_skip_block and a_txt.lower().startswith(SKIP_FROM.lower()):
            in_skip_block = True

        if in_skip_block:
            if a_txt.lower().startswith(SKIP_UNTIL_INCLUSIVE.lower()):
                in_skip_block = False
            continue

        if a_txt and "total" in a_txt.lower():
            continue

        if a_txt:
            start_new_site_template(a_txt)

        if b_txt and current_st is not None:
            ne_obj = {"network_element": b_txt}
            if c_txt:
                ne_obj["CNF_VNF"] = c_txt
            if d_txt:
                ne_obj["Subs"] = d_txt

            # all location columns for this row
            for loc in locations:
                loc_name = loc["name"]
                col_idx = loc["col_idx"]
                ne_obj[loc_name] = to_int_or_zero(ws.cell(row=r, column=col_idx).value)

            current_st["network_elements"].append(ne_obj)

    payload = {
        "locations": [loc["name"] for loc in locations],
        "Site-Templates": site_templates
    }

    out_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    print(f"Wrote JSON to {out_path} ({len(site_templates)} site-templates)")


if __name__ == "__main__":
    main()