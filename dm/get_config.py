#!/usr/bin/env python3
"""
Parse an XLSX sheet "Detailed Dimensioning- 172M" into JSON.

Adds filename parsing:
- version: extracted from ..._v<version>_...
- program: extracted from basename portion after "Assignment_"
  Example:
    USP_Evolution_CNF_VNF_Resources_2023TPA_v17.3_VM_AZ_Assignment_e2509.xlsx
    program = "e2509"
"""

import argparse
import json
import math
import os
import re
import sys
from typing import Any, Dict, List, Optional, Union

from openpyxl import load_workbook


def is_blank(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and v.strip() == "":
        return True
    return False


def clean_scalar(v: Any) -> Any:
    """Normalize cell values for JSON."""
    if is_blank(v):
        return None
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, float):
        if math.isfinite(v) and v.is_integer():
            return int(v)
        return v
    return v


def extract_version_from_filename(path: str) -> Optional[str]:
    """Extract version like '17.3' from ..._v17.3_..."""
    name = os.path.basename(path)
    m = re.search(r"_v([^_]+)_", name)
    return m.group(1) if m else None


def extract_program_from_filename(path: str) -> Optional[str]:
    """
    Extract program from basename after 'Assignment_'.

    Example:
      ..._Assignment_e2509.xlsx -> e2509
      ..._Assignment_e2509_anything.xlsx -> e2509 (only up to next underscore)
    """
    base = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r"Assignment_([^_]+)", base)
    return m.group(1) if m else None


def deprecate_json_path_for_xlsx(xlsx_path: str) -> str:
    base, _ = os.path.splitext(xlsx_path)
    return base + "-deprecate.json"


def load_deprecate_status_map(dep_path: str) -> Dict[str, Any]:
    """
    Loads the deprecate json and returns a mapping:
      config_heading (str) -> status

    Accepts either:
      - a list of objects: [{"config_heading": "...", "status": "..."}, ...]
      - or an object with a top-level list under a common key (best effort)
    """
    with open(dep_path, "r", encoding="utf-8") as f:
        payload: Union[List[Any], Dict[str, Any]] = json.load(f)

    items: Optional[List[Dict[str, Any]]] = None
    if isinstance(payload, list):
        items = payload
    elif isinstance(payload, dict):
        for v in payload.values():
            if isinstance(v, list):
                items = v
                break

    if items is None:
        raise ValueError(
            f"Deprecate JSON format not understood in {dep_path}. "
            "Expected a list or a dict containing a list."
        )

    status_map: Dict[str, Any] = {}
    for obj in items:
        if not isinstance(obj, dict):
            continue
        ch = obj.get("config_heading")
        if ch is None:
            continue
        ch = str(ch).strip()
        if not ch:
            continue
        status_map[ch] = obj.get("status")
    return status_map


def parse_xlsx(
    path: str,
    sheet_name: str,
    start_row: int = 22,
    deprecate_status_map: Optional[Dict[str, Any]] = None
) -> Dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f'Sheet "{sheet_name}" not found. Available: {wb.sheetnames}')

    ws = wb[sheet_name]

    configs: List[Dict[str, Any]] = []
    current: Optional[Dict[str, Any]] = None

    for r in range(start_row, ws.max_row + 1):
        b = clean_scalar(ws.cell(row=r, column=2).value)  # B
        c = clean_scalar(ws.cell(row=r, column=3).value)  # C
        d = clean_scalar(ws.cell(row=r, column=4).value)  # D

        e = clean_scalar(ws.cell(row=r, column=5).value)  # E
        f = clean_scalar(ws.cell(row=r, column=6).value)  # F
        g = clean_scalar(ws.cell(row=r, column=7).value)  # G
        h = clean_scalar(ws.cell(row=r, column=8).value)  # H
        i = clean_scalar(ws.cell(row=r, column=9).value)  # I
        j = clean_scalar(ws.cell(row=r, column=10).value) # J

        # New config starts when column B has an entry
        if not is_blank(b):
            status_val: Any = None
            if deprecate_status_map is not None:
                status_val = deprecate_status_map.get(str(b), "ERROR - not found")

            current = {
                "config_heading": b,
                "config_heading_row": r,
                "CNF_VNF": c,
                "subscription_count": d,
                "status": status_val,
                "items": []
            }
            configs.append(current)

            # Sometimes the same row might also have a sub-category in E..J
            if not is_blank(e):
                current["items"].append({
                    "pod_vm": e,
                    "quantity": f,
                    "vcpu": g,
                    "memory_G": h,
                    "root_disk_GB": i,
                    "cinder_gb": j,
                })
            continue

        if current is None:
            continue

        if is_blank(e):
            continue

        current["items"].append({
            "pod_vm": e,
            "quantity": f,
            "vcpu": g,
            "memory_G": h,
            "root_disk_GB": i,
            "cinder_gb": j,
        })

    return {
        "source_file": os.path.basename(path),
        "version": extract_version_from_filename(path),
        "program": extract_program_from_filename(path),
        "sheet": sheet_name,
        "start_row": start_row,
        "configs": configs
    }


def main():
    ap = argparse.ArgumentParser(description="Parse dimensioning XLSX to JSON.")
    ap.add_argument("xlsx", help="Input .xlsx filename")
    ap.add_argument(
        "-o", "--output",
        help="Output JSON filename (default: <input>.json)",
        default=None
    )
    ap.add_argument(
        "--sheet",
        default="Detailed Dimensioning- 172M",
        help='Sheet name (default: "Detailed Dimensioning- 172M")'
    )
    ap.add_argument(
        "--start-row",
        type=int,
        default=22,
        help="1-based Excel row to start reading (default: 22)"
    )
    ap.add_argument(
        "--read-deprecate-json",
        action="store_true",
        help='If set, require and use "<xlsx base>-deprecate.json" to populate config status.'
    )
    args = ap.parse_args()

    out = args.output
    if out is None:
        base, _ = os.path.splitext(args.xlsx)
        out = base + ".json"

    deprecate_status_map: Optional[Dict[str, Any]] = None
    dep_path: Optional[str] = None

    if args.read_deprecate_json:
        dep_path = deprecate_json_path_for_xlsx(args.xlsx)
        if not os.path.isfile(dep_path):
            print(f"ERROR: Deprecate JSON file not found: {dep_path}", file=sys.stderr)
            sys.exit(1)
        try:
            deprecate_status_map = load_deprecate_status_map(dep_path)
        except json.JSONDecodeError as e:
            print(f"ERROR: Failed to parse deprecate JSON ({dep_path}): {e}", file=sys.stderr)
            sys.exit(1)
        except Exception as e:
            print(f"ERROR: {e}", file=sys.stderr)
            sys.exit(1)

    data = parse_xlsx(
        args.xlsx,
        args.sheet,
        args.start_row,
        deprecate_status_map=deprecate_status_map
    )

    if args.read_deprecate_json and dep_path is not None:
        data["deprecate_json_file"] = os.path.basename(dep_path)

    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    print(f"Wrote {out}")


if __name__ == "__main__":
    main()