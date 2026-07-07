#!/usr/bin/env python3
import sys
import json
from pathlib import Path
from collections import OrderedDict
import openpyxl

SHEET_NAME = "CNF VNF Counts - 172M"

# Site names are on row 23, starting from column E
LOCATION_ROW = 23
LOCATION_START_COL = 5  # E
LOCATION_STOP_AFTER_EMPTY = 3  # stop after N consecutive empty cells

# Data starts here
START_ROW = 25

# Column indices
COL_SITE_TEMPLATE = 1    # A
COL_NETWORK_ELEMENT = 2  # B
COL_CNF_VNF = 3          # C
COL_SUBS = 4             # D

SKIP_FROM = "3rd Party Onboarding"
SKIP_UNTIL_INCLUSIVE = "Infrastructure"
STOP_CONTAINS = "Additional (Nokia)"


def normalize_cell_text(v) -> str:
    s = str(v).strip()
    s = s.replace("\r", " ").replace("\n", " ")
    s = " ".join(s.split())
    return s


def cell_to_string(v) -> str:
    if v is None:
        return ""
    return normalize_cell_text(v)


def excel_col_name(n: int) -> str:
    name = ""
    while n:
        n, r = divmod(n - 1, 26)
        name = chr(65 + r) + name
    return name


def count_value(v):
    """Convert a cell value into a count.
       - None / empty -> 0
       - 'DNP' (any case) -> 0
       - int/float -> numeric (int if whole)
       - numeric string -> numeric (int if whole)
       - anything else -> 0
    """
    if v is None or isinstance(v, bool):
        return 0
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else v
    s = normalize_cell_text(v)
    if s == "" or s.upper() == "DNP":
        return 0
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except ValueError:
        return 0


def read_sites(ws):
    """Read site names from LOCATION_ROW starting at LOCATION_START_COL.
       Returns list of (col_index, site_name)."""
    sites = []
    empty_streak = 0
    col = LOCATION_START_COL
    while True:
        v = ws.cell(row=LOCATION_ROW, column=col).value
        s = cell_to_string(v)
        if s == "":
            empty_streak += 1
            if empty_streak >= LOCATION_STOP_AFTER_EMPTY:
                break
        else:
            empty_streak = 0
            sites.append((col, s))
        col += 1
    return sites


def main():
    if len(sys.argv) != 2:
        print(f"Usage: {Path(sys.argv[0]).name} <workbook.xlsx>", file=sys.stderr)
        sys.exit(2)

    in_path = Path(sys.argv[1])
    wb = openpyxl.load_workbook(in_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"Sheet '{SHEET_NAME}' not found in {in_path}", file=sys.stderr)
        sys.exit(1)
    ws = wb[SHEET_NAME]

    sites = read_sites(ws)
    if not sites:
        print("No sites found on row 23", file=sys.stderr)
        sys.exit(1)

    # Preserve site order
    site_data = OrderedDict()   # site_name -> OrderedDict(template -> list of NE dicts)
    for _, sname in sites:
        site_data[sname] = OrderedDict()

    row = START_ROW
    max_row = ws.max_row
    current_template = ""
    skipping = False

    while row <= max_row:
        template_val = cell_to_string(ws.cell(row=row, column=COL_SITE_TEMPLATE).value)
        ne_val       = cell_to_string(ws.cell(row=row, column=COL_NETWORK_ELEMENT).value)
        cnf_val      = cell_to_string(ws.cell(row=row, column=COL_CNF_VNF).value)
        subs_val     = cell_to_string(ws.cell(row=row, column=COL_SUBS).value)

        # Update current template if this row starts a new one
        if template_val:
            current_template = template_val

        # Stop condition
        combined = " ".join([template_val, ne_val]).strip()
        if STOP_CONTAINS and STOP_CONTAINS in combined:
            break

        # Skip section: from SKIP_FROM up to and including SKIP_UNTIL_INCLUSIVE
        if SKIP_FROM and (template_val == SKIP_FROM or ne_val == SKIP_FROM):
            skipping = True
        if skipping:
            if SKIP_UNTIL_INCLUSIVE and (
                template_val == SKIP_UNTIL_INCLUSIVE or ne_val == SKIP_UNTIL_INCLUSIVE
            ):
                skipping = False
            row += 1
            continue

        # Rows without a network element are section headers; skip them
        if not ne_val:
            row += 1
            continue

        # For every site column, add this network element (including 0-count)
        for col, sname in sites:
            cnt = count_value(ws.cell(row=row, column=col).value)
            tkey = current_template if current_template else ""
            if tkey not in site_data[sname]:
                site_data[sname][tkey] = []
            site_data[sname][tkey].append({
                "network_element": ne_val,
                "CNF_VNF": cnf_val,
                "Subs": subs_val,
                "line": row,
                "count": cnt,
            })

        row += 1

    # Build final JSON structure
    output = []
    for sname, templates in site_data.items():
        site_entry = {
            "site": sname,
            "Site-Templates": []
        }
        for tname, nes in templates.items():
            site_entry["Site-Templates"].append({
                "site-template": tname,
                "network_elements": nes,
            })
        output.append(site_entry)

    # Write to <input-basename>-siteinfo.json in the same directory as the input file
    out_path = in_path.with_name(f"{in_path.stem}-siteinfo.json")
    with out_path.open("w", encoding="utf-8") as f:
        json.dump(output, f, indent=2)

    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()