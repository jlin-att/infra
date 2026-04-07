#!/usr/bin/env python3

# take a DM file, check on the color and use color to determine if the row is active or deprecated. Write a JSON file with the results.


import argparse
import json
import os
from collections import Counter
from openpyxl import load_workbook

SHEET_NAME = "Detailed Dimensioning- 172M"
START_ROW = 22  # start scanning from Excel row 22

def color_to_text(color):
    if color is None:
        return "none"

    ctype = getattr(color, "type", None)
    if ctype == "rgb":
        return f"rgb:{color.rgb}"
    if ctype == "theme":
        tint = getattr(color, "tint", None)
        return f"theme:{color.theme}" + (f":tint={tint}" if tint is not None else "")
    if ctype == "indexed":
        return f"indexed:{color.indexed}"

    rgb = getattr(color, "rgb", None)
    if rgb:
        return f"rgb:{rgb}"
    return "unknown"

def theme_tint_key(color):
    """Return (theme, tint) if this is a theme color, else None."""
    if color is None:
        return None
    if getattr(color, "type", None) == "theme":
        return (getattr(color, "theme", None), getattr(color, "tint", None))
    return None

def status_from_color(color):
    """
    ACTIVE when:
      - no color / can't detect color (None or unknown type), OR
      - theme==1 and tint==0.0
    Otherwise DEPRECATED.
    """
    if color is None:
        return "active"

    ctype = getattr(color, "type", None)
    if ctype is None:
        return "active"

    if ctype == "theme":
        theme = getattr(color, "theme", None)
        tint = getattr(color, "tint", None)
        if theme == 1 and tint == 0.0:
            return "active"
        return "deprecated"

    if ctype in ("rgb", "indexed"):
        return "deprecated"

    return "active"

def main():
    ap = argparse.ArgumentParser(
        description="Extract non-empty cells from column B and write JSON to a file; print theme/tint summary to stdout."
    )
    ap.add_argument("xlsx", help="Path to .xlsx file")
    ap.add_argument(
        "-o", "--out",
        help="Output JSON filename (default: <input>-deprecate.json)",
        default=None
    )
    ap.add_argument("--max-rows", type=int, default=None, help="Optional limit on rows scanned")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f'Sheet not found: "{SHEET_NAME}". Available: {wb.sheetnames}')

    ws = wb[SHEET_NAME]

    max_row = ws.max_row
    if args.max_rows is not None:
        max_row = min(max_row, args.max_rows)

    theme_tint_counts = Counter()
    records = []

    for r in range(START_ROW, max_row + 1):
        cell = ws.cell(row=r, column=2)  # column B
        val = cell.value

        if val is None:
            continue
        if isinstance(val, str) and val.strip() == "":
            continue

        color = getattr(cell.font, "color", None)

        key = theme_tint_key(color)
        if key is not None:
            theme_tint_counts[key] += 1

        records.append({
            "row": r,
            "config_heading": val,
            "font_color": color_to_text(color),
            "status": status_from_color(color),
        })

    out = args.out
    if out is None:
        base, _ = os.path.splitext(args.xlsx)
        out = base + "-deprecate.json"
    # Write JSON output file (no theme/tint summary in the file)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2, default=str)

    # Summary still goes to stdout (screen)
    print("Theme/Tint combinations detected (theme, tint) -> count")
    if not theme_tint_counts:
        print("  (none)")
    else:
        def sort_key(item):
            (theme, tint), _cnt = item
            tint_sort = tint if tint is not None else 0.0
            return (theme if theme is not None else -1, tint_sort)

        for (theme, tint), cnt in sorted(theme_tint_counts.items(), key=sort_key):
            print(f"  ({theme}, {tint}) -> {cnt}")

if __name__ == "__main__":
    main()