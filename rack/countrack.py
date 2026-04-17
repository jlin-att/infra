#!/usr/bin/env python3
"""
Process an Excel (.xlsx) file provided via command line.
- Finds sheets containing '+' in their name.
- Locates 'Rack <1-15>' cells within the first 5 rows (spanning 2 merged cols).
- For each Rack, reads ALL rows below the header in the SECOND column only.
- Replaces line feeds with spaces in cell values.
- Determines U-Size (1, 2, or 3) based on whether a cell is part of a
  vertically merged range spanning 2 or 3 rows in the same column.
- Counts value occurrences (by value AND u-size) per Rack and per Sheet.
- Outputs results in CSV format.

Usage: python script.py <path_to_excel_file>
"""

import sys
import os
import re
import csv
import openpyxl
from collections import defaultdict


# Matches "Rack <N>" where N is 1-15 (space optional, case-insensitive)
RACK_PATTERN = re.compile(r"^Rack\s*(1[0-5]|[1-9])$", re.IGNORECASE)


def normalize_value(value) -> str:
    """Strip and replace any line feed (\\n or \\r\\n) with a single space."""
    return re.sub(r"[\r\n]+", " ", str(value).strip())


def get_merged_2col_ranges(sheet) -> list:
    """Return all merged ranges that span exactly 2 columns."""
    return [
        mr for mr in sheet.merged_cells.ranges
        if mr.max_col - mr.min_col == 1
    ]


def build_vertical_merge_lookup(sheet, col: int) -> dict:
    """
    Build a lookup dict for vertically merged cells in a specific column.
    Returns: {min_row: span_size} for merged ranges where min_col == max_col == col
    and the span is exactly 2 or 3 rows.
    """
    lookup = {}
    for mr in sheet.merged_cells.ranges:
        if mr.min_col == col and mr.max_col == col:
            span = mr.max_row - mr.min_row + 1
            if span in (2, 3):
                lookup[mr.min_row] = span
    return lookup


def find_rack_headers(sheet) -> list:
    """
    Scan the first 5 rows for cells matching 'Rack <1-15>'
    that are part of a 2-column horizontal merge.
    Returns a list of dicts: {rack_label, header_row, col_start, col_end}
    """
    merged_2col = get_merged_2col_ranges(sheet)
    rack_headers = []

    for row in sheet.iter_rows(min_row=1, max_row=5):
        for cell in row:
            if cell.value is None:
                continue
            value_str = normalize_value(cell.value)
            if RACK_PATTERN.match(value_str):
                for mr in merged_2col:
                    if mr.min_row == cell.row and mr.min_col == cell.column:
                        rack_headers.append({
                            "rack_label": value_str,
                            "header_row": cell.row,
                            "col_start":  mr.min_col,
                            "col_end":    mr.max_col,
                        })
                        break

    return rack_headers


def count_values_in_second_col(sheet, rack_info: dict) -> dict:
    """
    Read ALL rows below the rack header, SECOND column of the merge only.
    Normalize line feeds to spaces. Skip empty cells.
    Determine U-Size:
      - 3 if the cell is part of a vertical merge spanning 3 rows
      - 2 if the cell is part of a vertical merge spanning 2 rows
      - 1 otherwise (single cell, no vertical merge)
    Returns a dict: {(value, u_size): count}
    """
    data_col  = rack_info["col_end"]
    start_row = rack_info["header_row"] + 1
    counts    = defaultdict(int)

    # Build lookup of vertically merged cells (2- or 3-row spans) in this column
    vmerge_lookup = build_vertical_merge_lookup(sheet, data_col)

    row = start_row
    while row <= sheet.max_row:
        raw = sheet.cell(row=row, column=data_col).value
        if raw is not None:
            val = normalize_value(raw)
            if row in vmerge_lookup:
                u_size = vmerge_lookup[row]
                counts[(val, u_size)] += 1
                row += u_size  # skip the merged rows
                continue
            else:
                u_size = 1
                counts[(val, u_size)] += 1
        row += 1

    return dict(counts)


def process_file(filepath: str) -> None:
    # --- Validate ---
    if not os.path.exists(filepath):
        print(f"[ERROR] File not found: {filepath}", file=sys.stderr)
        sys.exit(1)

    if not filepath.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        print(f"[ERROR] Not a valid Excel file: {filepath}", file=sys.stderr)
        sys.exit(1)

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        print(f"[ERROR] Could not open workbook: {e}", file=sys.stderr)
        sys.exit(1)

    matched_sheets = [name for name in wb.sheetnames if "+" in name]

    if not matched_sheets:
        print("No sheets found containing '+'.", file=sys.stderr)
        wb.close()
        return

    # sheet_summary[sheet_name][rack_label][(value, u_size)] = count
    sheet_summary = {}

    for sheet_name in matched_sheets:
        sheet = wb[sheet_name]
        rack_headers = find_rack_headers(sheet)
        rack_summary = {}

        for rack_info in rack_headers:
            counts = count_values_in_second_col(sheet, rack_info)
            if counts:
                rack_summary[rack_info["rack_label"]] = counts

        if rack_summary:
            sheet_summary[sheet_name] = rack_summary

    wb.close()

    # ------------------------------------------------------------------ #
    #  CSV OUTPUT                                                          #
    # ------------------------------------------------------------------ #
    writer = csv.writer(sys.stdout)

    for sheet_name, racks in sheet_summary.items():

        # Collect all unique (value, u_size) keys across racks
        all_keys  = sorted({k for counts in racks.values() for k in counts})
        rack_labels = sorted(racks.keys(),
                             key=lambda r: int(re.search(r"\d+", r).group()))

        # --- Per-Rack section ---
        writer.writerow([f"Sheet: {sheet_name}"])
        writer.writerow(["Value", "U-Size"] + rack_labels)

        for val, u_size in all_keys:
            writer.writerow(
                [val, u_size] +
                [racks[rack].get((val, u_size), 0) for rack in rack_labels]
            )

        # Totals row per rack
        writer.writerow(
            ["TOTAL", ""] +
            [sum(racks[rack].values()) for rack in rack_labels]
        )

        # Blank separator
        writer.writerow([])

        # --- Per-Sheet aggregate ---
        writer.writerow([f"Sheet Totals: {sheet_name}"])
        writer.writerow(["Value", "U-Size", "Count"])

        sheet_totals = defaultdict(int)
        for counts in racks.values():
            for key, cnt in counts.items():
                sheet_totals[key] += cnt

        for (val, u_size) in sorted(sheet_totals):
            writer.writerow([val, u_size, sheet_totals[(val, u_size)]])

        writer.writerow(["TOTAL", "", sum(sheet_totals.values())])

        # Blank separator between sheets
        writer.writerow([])


def main():
    if len(sys.argv) != 2:
        print("Usage: python script.py <path_to_excel_file>", file=sys.stderr)
        sys.exit(1)

    process_file(sys.argv[1])


if __name__ == "__main__":
    main()
